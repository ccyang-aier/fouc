"""Generate actual PDF/OOXML fixtures with repeatable headings, tables and pixels."""
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches as WordInches
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image as PdfImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / ".cache" / "documents"
TITLE = "Fouc Knowledge Report"
HEADING = "Quarterly Capacity"
CAPTION = "Figure 1. Capacity chart."
ROWS = [["Item", "Count", "Status"], ["Alpha", "12", "Ready"], ["Beta", "7", "Review"]]


def main():
    SAMPLES.mkdir(parents=True, exist_ok=True)
    image_path = SAMPLES / "capacity.png"
    image = Image.new("RGB", (480, 260), "#f3f5f7")
    draw = ImageDraw.Draw(image)
    draw.line((60, 30, 60, 215, 440, 215), fill="#334155", width=3)
    draw.rectangle((120, 65, 220, 212), fill="#146b5f")
    draw.rectangle((280, 125, 380, 212), fill="#8b5e3c")
    draw.text((150, 225), "Alpha", fill="#162f2a")
    draw.text((315, 225), "Beta", fill="#162f2a")
    image.save(image_path)

    styles = getSampleStyleSheet()
    table = Table(ROWS, colWidths=[170, 120, 150])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5eee9")),
                              ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#82978c")),
                              ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                              ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    SimpleDocTemplate(str(SAMPLES / "report.pdf"), title=TITLE, author="Fouc development verification").build([
        Paragraph(TITLE, styles["Title"]), Spacer(1, 14), Paragraph(HEADING, styles["Heading1"]),
        Paragraph("This generated document verifies real structured conversion.", styles["BodyText"]),
        Spacer(1, 18), table, Spacer(1, 24), PdfImage(str(image_path), width=360, height=195),
        Spacer(1, 10), Paragraph(CAPTION, styles["Caption"] if "Caption" in styles else styles["BodyText"]),
    ])
    blank_pdf = Canvas(str(SAMPLES / "empty.pdf"))
    blank_pdf.showPage()
    blank_pdf.save()

    word = Document()
    word.add_heading(TITLE, 0)
    word.add_heading(HEADING, 1)
    word.add_paragraph("This generated document verifies real structured conversion.")
    word_table = word.add_table(rows=0, cols=3)
    word_table.style = "Table Grid"
    for row in ROWS:
        for cell, text in zip(word_table.add_row().cells, row):
            cell.text = text
    word.add_picture(str(image_path), width=WordInches(4.5))
    word.add_paragraph(CAPTION, style="Caption")
    word.save(SAMPLES / "report.docx")
    Document().save(SAMPLES / "empty.docx")

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = TITLE
    heading = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(8), Inches(0.6)).text_frame
    heading.text = HEADING
    heading.paragraphs[0].font.size = Pt(24)
    deck_table = slide.shapes.add_table(3, 3, Inches(0.5), Inches(2.1), Inches(4.3), Inches(2)).table
    for row_index, row in enumerate(ROWS):
        for column_index, text in enumerate(row):
            deck_table.cell(row_index, column_index).text = text
    slide.shapes.add_picture(str(image_path), Inches(5.1), Inches(2.1), width=Inches(4.3))
    slide.shapes.add_textbox(Inches(5.1), Inches(4.55), Inches(4.3), Inches(0.6)).text = CAPTION
    deck.save(SAMPLES / "report.pptx")
    Presentation().save(SAMPLES / "empty.pptx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = HEADING
    sheet["A1"] = TITLE
    sheet["A1"].font = Font(size=20, bold=True)
    sheet.merge_cells("A1:C1")
    for row_index, row in enumerate(ROWS, 3):
        for column_index, text in enumerate(row, 1):
            cell = sheet.cell(row_index, column_index, int(text) if text.isdecimal() else text)
            if row_index == 3:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E5EEE9")
    for column in ("A", "B", "C"):
        sheet.column_dimensions[column].width = 20
    sheet.add_image(SheetImage(str(image_path)), "E3")
    sheet["E18"] = CAPTION
    workbook.save(SAMPLES / "report.xlsx")
    Workbook().save(SAMPLES / "empty.xlsx")

    evidence = {"provenance": "Locally generated using reportlab, python-docx, python-pptx, openpyxl and Pillow; no external documents or credentials.", "files": []}
    for path in sorted(SAMPLES.iterdir()):
        if path.suffix not in (".pdf", ".docx", ".pptx", ".xlsx", ".png"):
            continue
        evidence["files"].append({"name": path.name, "bytes": path.stat().st_size,
                                  "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    (SAMPLES / "samples.json").write_text(json.dumps(evidence, indent=2) + "\n")
    print(json.dumps(evidence, indent=2))


if __name__ == "__main__":
    main()
