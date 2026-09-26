import base64
import hashlib
import io
import multiprocessing
import socket
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from PIL import Image
from reportlab.pdfgen.canvas import Canvas

from media_worker.contracts import Attachment, DOCUMENT_TYPES, ProcessorOutput
from media_worker.config import Settings
from media_worker.docling import offline_document_io, parse_document
from media_worker.docling_config import DoclingSettings
from media_worker.document_limits import preflight_document
from media_worker.errors import WorkerError
from media_worker.processors import ProcessorFailure, ProcessorInput, ProcessorSpec, run_processor


def source(path):
    mime = next(mime for mime, extension in DOCUMENT_TYPES.items() if extension == path.suffix)
    return ProcessorInput(path, mime, None, "cpu", 0, "auto", 2)


def word_fixture(directory: Path, pictures=1):
    word = Document()
    word.add_heading("Real document heading", 1)
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Product", "Quantity"
    table.cell(1, 0).text, table.cell(1, 1).text = "Orchard", "17"
    for index in range(pictures):
        encoded = io.BytesIO()
        Image.new("RGB", (20, 10), (index, 70, 50)).save(encoded, "PNG")
        encoded.seek(0)
        word.add_picture(encoded)
    path = directory / "actual.docx"
    word.save(path)
    return source(path)


def test_configuration_rejects_unbounded_document_limits(tmp_path):
    config = DoclingSettings.from_environment({"MEDIA_WORKER_DOCLING_ARTIFACTS_PATH": str(tmp_path),
                                               "MEDIA_WORKER_DOCLING_TABLE_MODE": "accurate"})
    assert not config.is_prepared() and "accurate" in config.model_files()[1][2][0]
    for values in ({"MEDIA_WORKER_DOCLING_MAX_PAGES": "0"}, {"MEDIA_WORKER_DOCLING_TABLE_MODE": "remote"},
                   {"MEDIA_WORKER_DOCLING_MAX_IMAGES": "33"}, {"MEDIA_WORKER_DOCLING_MAX_IMAGE_PIXELS": "4000001"},
                   {"MEDIA_WORKER_DOCLING_MAX_TOTAL_IMAGE_BYTES": "1048577"}):
        with pytest.raises(ValueError):
            DoclingSettings.from_environment(values)


def test_document_conversion_refuses_network_and_restores_socket_api():
    connect = socket.socket.connect
    with offline_document_io():
        with pytest.raises(PermissionError):
            socket.create_connection(("127.0.0.1", 9))
        with pytest.raises(PermissionError):
            socket.getaddrinfo("not-requested.example", 80)
    assert socket.socket.connect is connect


def test_real_docx_heading_table_pixels_and_separate_attachment(tmp_path):
    output = ProcessorOutput.model_validate(parse_document(word_fixture(tmp_path), config=DoclingSettings(artifacts_path=tmp_path)))
    assert "## Real document heading" in output.derived.markdown
    assert "Orchard" in output.derived.markdown and "17" in output.derived.markdown and "|" in output.derived.markdown
    assert len(output.attachments) == 1 and "data:" not in output.derived.markdown
    attachment = output.attachments[0]
    data = base64.b64decode(attachment.dataBase64)
    with Image.open(io.BytesIO(data)) as pixels:
        assert pixels.size == (20, 10) and pixels.getpixel((0, 0)) == (0, 70, 50)
    assert hashlib.sha256(data).hexdigest() == attachment.sha256
    assert f"asset:{attachment.sha256}" in output.derived.markdown


@pytest.mark.parametrize("override", [{"max_images": 0}, {"max_image_pixels": 199},
                                       {"max_image_bytes": 20}, {"max_total_image_bytes": 20},
                                       {"max_uncompressed_bytes": 1024}, {"max_zip_entries": 1}])
def test_real_document_limits_fail_without_silently_dropping_images(tmp_path, override):
    with pytest.raises(ProcessorFailure, match="document_limit_exceeded"):
        parse_document(word_fixture(tmp_path), config=DoclingSettings(**override))


def test_real_two_images_cannot_exceed_aggregate_limit(tmp_path):
    document = word_fixture(tmp_path, pictures=2)
    with pytest.raises(ProcessorFailure, match="document_limit_exceeded"):
        parse_document(document, config=DoclingSettings(max_total_image_bytes=100))
    with pytest.raises(ProcessorFailure, match="document_limit_exceeded"):
        parse_document(document, config=DoclingSettings(max_images=1))


@pytest.mark.parametrize("kind", ["path", "entity", "entity_utf16"])
def test_ooxml_rejects_traversal_and_entities(tmp_path, kind):
    document = word_fixture(tmp_path)
    with ZipFile(document.source, "a") as archive:
        content = '<!DOCTYPE x [<!ENTITY leak SYSTEM "file:///not-read">]><x/>'
        archive.writestr("../unsafe.txt" if kind == "path" else "unsafe.xml",
                         "unsafe" if kind == "path" else content.encode("utf-16" if kind == "entity_utf16" else "utf-8"))
    with pytest.raises(ProcessorFailure, match="invalid_document"):
        preflight_document(document, DoclingSettings())


def test_pdf_page_and_render_pixel_limits_before_inference(tmp_path):
    path = tmp_path / "pages.pdf"
    pdf = Canvas(str(path))
    for _ in range(2):
        pdf.drawString(50, 50, "A real PDF page")
        pdf.showPage()
    pdf.save()
    for config in (DoclingSettings(max_pages=1), DoclingSettings(max_page_pixels=10_000)):
        with pytest.raises(ProcessorFailure, match="document_limit_exceeded"):
            preflight_document(source(path), config)
    with pytest.raises(ProcessorFailure, match="model_unavailable"):
        parse_document(source(path), config=DoclingSettings(artifacts_path=tmp_path))


def test_corrupt_prepared_pdf_model_is_non_retryable_configuration_failure(tmp_path):
    path = tmp_path / "text.pdf"
    pdf = Canvas(str(path))
    pdf.drawString(50, 50, "A real text layer")
    pdf.showPage()
    pdf.save()
    config = DoclingSettings(artifacts_path=tmp_path / "models")
    for repository, _revision, files in config.model_files():
        for filename in files:
            artifact = config.artifacts_path / repository.replace("/", "--") / filename
            artifact.parent.mkdir(parents=True, exist_ok=True)
            artifact.write_bytes(b"deliberately corrupted model, never used as recognition evidence")
    with pytest.raises(ProcessorFailure, match="model_unavailable"):
        parse_document(source(path), config=config)


def test_office_conversion_does_not_discover_external_applications(tmp_path, monkeypatch):
    from docling.backend.docx.drawingml import utils

    def unexpected_discovery():
        pytest.fail("Document parser must not discover/launch LibreOffice")

    monkeypatch.setattr(utils, "get_libreoffice_cmd", unexpected_discovery)
    result = parse_document(word_fixture(tmp_path), config=DoclingSettings())
    assert result["attachments"]
    assert utils.get_libreoffice_cmd is unexpected_discovery


@pytest.mark.parametrize("extension", [".pdf", ".docx", ".pptx", ".xlsx"])
def test_corrupt_document_error_is_explicit(tmp_path, extension):
    path = tmp_path / f"corrupt{extension}"
    path.write_bytes(b"This is deliberately not a document container.")
    with pytest.raises(ProcessorFailure, match="invalid_document"):
        parse_document(source(path), config=DoclingSettings())


def test_empty_word_document_is_not_a_ready_empty_result(tmp_path):
    path = tmp_path / "empty.docx"
    Document().save(path)
    with pytest.raises(ProcessorFailure, match="empty_document"):
        parse_document(source(path), config=DoclingSettings())


def test_attachment_contract_checks_integrity_shape_and_references(tmp_path):
    output = parse_document(word_fixture(tmp_path), config=DoclingSettings())
    attachment = output["attachments"][0]
    for change in ({"sha256": "a" * 64}, {"size": 1}, {"width": 1}, {"mime": "image/jpeg"},
                   {"dataBase64": "invalid"}, {"width": 4_000_000, "height": 2}):
        with pytest.raises(ValueError):
            Attachment.model_validate({**attachment, **change})
    with pytest.raises(ValueError):
        ProcessorOutput.model_validate({**output, "attachments": [attachment, attachment]})
    with pytest.raises(ValueError):
        ProcessorOutput.model_validate({**output, "attachments": []})
    with pytest.raises(ValueError):
        ProcessorOutput.model_validate({**output, "derived": {"status": "ready", "markdown": "No image references"}})


async def test_real_docling_child_obeys_aggregate_response_limit(tmp_path):
    document = word_fixture(tmp_path)
    word = Document(document.source)
    word.add_paragraph("Bounded document response content. " * 100)
    word.save(document.source)
    settings = Settings(token="test-only-" + "x" * 32, download_origins=("http://storage.test",), max_response_bytes=1024)
    spec = ProcessorSpec("docling", "media_worker.docling:parse_document", {"config": DoclingSettings()})
    with pytest.raises(WorkerError) as error:
        await run_processor(spec, document, settings)
    assert error.value.code == "result_too_large" and error.value.status == 413 and not error.value.retryable
    assert not multiprocessing.active_children()
